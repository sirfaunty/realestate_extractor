"""Investor Report (.docx) for the Partnership Dashboard — pure Python.

Replaces the former Node generator (generate_report.js), so the export needs
nothing beyond python-docx (already required by Deliverables) on a laptop or in
the container.

    build_investor_report(dashboard_dict, out_path)

`dashboard_dict` is DashboardResult.to_dict(), optionally with
'_primary_scenario' naming the featured scenario.
"""

from __future__ import annotations

import datetime

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x3D, 0x6B)
PRIMARY = RGBColor(0x18, 0x5F, 0xA5)
GRAY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xCC, 0x00, 0x00)
AMBER = RGBColor(0xCC, 0x77, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FILL_HEADER = '0B3D6B'
FILL_ALT = 'F5F5F5'
FILL_TOTAL = 'E6F1FB'
FONT = 'Arial'


# ─── formatting ────────────────────────────────────────────────────

def fmt_usd(v):
    if v is None:
        return '—'
    a, sign = abs(v), '-' if v < 0 else ''
    if a >= 1e6:
        return f'{sign}${a / 1e6:.2f}M'
    if a >= 1e3:
        return f'{sign}${round(a):,}'
    return f'{sign}${a:.0f}'


def fmt_pct(v):
    return '—' if v is None else f'{v * 100:.2f}%'


def fmt_x(v):
    return '—' if v is None else f'{v:.2f}x'


def fmt_k(v):
    return '—' if v is None else f'${round(v / 1000):,}K'


# ─── docx helpers ──────────────────────────────────────────────────

def _run(par, text, size=9, bold=False, color=None, italic=False):
    r = par.add_run(str(text))
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def _para(doc_or_cell, text='', size=9, align=None, after=0, **kw):
    p = doc_or_cell.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    if text != '':
        _run(p, text, size=size, **kw)
    return p


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _set_cell(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT, fill=None,
              bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    _run(p, text, size=size, bold=bold, color=color)
    if fill:
        _shade(cell, fill)


def _table(doc, n_cols, widths_in):
    t = doc.add_table(rows=0, cols=n_cols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths_in):
        col.width = Inches(w)
    return t


def _add_row(t, widths_in):
    cells = t.add_row().cells
    for c, w in zip(cells, widths_in):
        c.width = Inches(w)
    return cells


def _header_row(t, widths, labels, first_align=WD_ALIGN_PARAGRAPH.LEFT):
    cells = _add_row(t, widths)
    for i, (c, h) in enumerate(zip(cells, labels)):
        _set_cell(c, h, align=first_align if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT,
                  fill=FILL_HEADER, bold=True, color=WHITE)


def _kv_table(doc, rows):
    widths = [4.6, 3.1]
    t = _table(doc, 2, widths)
    for i, (k, v) in enumerate(rows):
        cells = _add_row(t, widths)
        if k == '' and v == '':
            _set_cell(cells[0], ' ')
            _set_cell(cells[1], ' ')
            continue
        fill = FILL_ALT if i % 2 == 0 else None
        _set_cell(cells[0], k, fill=fill)
        _set_cell(cells[1], v, align=WD_ALIGN_PARAGRAPH.RIGHT, fill=fill, bold=True)
    return t


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    _run(p, text, size=16, bold=True, color=NAVY)
    return p


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _bottom_border(par, color='185FA5', size=6):
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', str(size)), ('w:space', '1'), ('w:color', color)):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def _page_number_field(par):
    r = par.add_run()
    r.font.name = FONT
    r.font.size = Pt(8)
    r.font.color.rgb = GRAY
    for tag, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
        r._r.append(el)


# ─── sections ──────────────────────────────────────────────────────

def _cover(doc, data, sc, date_str):
    for _ in range(8):
        _para(doc, after=6)
    _para(doc, data.get('entity_name') or 'Partnership', size=24, bold=True,
          color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    _para(doc, 'Partnership Investment Report', size=16, color=PRIMARY,
          align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    p = _para(doc, date_str, size=11, color=GRAY,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
    _bottom_border(p)
    _para(doc, after=20)
    _para(doc, f"TIF Scenario: {sc.get('tif_label', '')}", size=12, color=GRAY,
          align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    src = 'Live Proforma Engine' if sc.get('proforma_source') == 'live' else 'Default Assumptions'
    _para(doc, f'Data Source: {src}', size=10, color=GRAY,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    _page_break(doc)


def _deal_summary(doc, sc):
    dm, pf, debt = sc['decision_metrics'], sc['proforma'], sc['debt']
    _heading(doc, 'Deal Summary')
    _kv_table(doc, [
        ('Acquisition Cost Basis', fmt_usd(pf.get('acquisition_cost_basis'))),
        ('Total Equity', fmt_usd(pf.get('initial_equity'))),
        ('Hold Period', f"{pf.get('hold_years')} Years"),
        ('Exit Cap Rate', fmt_pct(pf.get('exit_cap_rate'))),
        ('Gross Sale Price', fmt_usd(pf.get('gross_sale_price'))),
        ('Net Sale Proceeds', fmt_usd(pf.get('net_sale_proceeds'))),
        ('', ''),
        ('Deal IRR', fmt_pct(dm.get('deal_irr'))),
        ('Deal Equity Multiple', fmt_x(dm.get('deal_em'))),
        ('Total Distributions (Operations)', fmt_usd(dm.get('total_distributions'))),
        ('Total Surplus Note Payments', fmt_usd(dm.get('total_surplus_note'))),
        ('Total MIP', fmt_usd(dm.get('total_mip'))),
        ('', ''),
        ('Senior Loan Balance', fmt_usd(debt.get('current_balance'))),
        ('Interest Rate', fmt_pct(debt.get('rate'))),
        ('Annual Debt Service', fmt_usd(debt.get('annual_debt_service'))),
        ('Min DSCR', fmt_x(dm.get('min_dscr'))),
        ('Avg DSCR', fmt_x(dm.get('avg_dscr'))),
    ])
    _para(doc, after=10)


def _partner_returns(doc, sc):
    _heading(doc, 'Partner Returns')
    widths = [2.0, 1.15, 1.15, 1.15, 1.15, 1.15]
    t = _table(doc, 6, widths)
    _header_row(t, widths, ['Partner', 'Equity', 'EM', 'IRR', 'Avg CoC', 'Unpaid Pref'])
    for i, p in enumerate(sc.get('partners', [])):
        fill = FILL_ALT if i % 2 == 0 else None
        vals = [p.get('name'), fmt_usd(p.get('initial_equity')), fmt_x(p.get('equity_multiple')),
                fmt_pct(p.get('irr')), fmt_pct(p.get('avg_cash_on_cash')), fmt_usd(p.get('unpaid_pref'))]
        for j, (c, v) in enumerate(zip(_add_row(t, widths), vals)):
            _set_cell(c, v, align=WD_ALIGN_PARAGRAPH.RIGHT if j else WD_ALIGN_PARAGRAPH.LEFT,
                      fill=fill, bold=(j == 0))
    _para(doc, after=6)
    for p in sc.get('partners', []):
        par = _para(doc, after=3)
        _run(par, f"{p.get('name')}: ", bold=True)
        _run(par, f"Accrued Pref {fmt_usd(p.get('accrued_pref'))} | Paid {fmt_usd(p.get('paid_pref'))}"
                  f" | Unpaid {fmt_usd(p.get('unpaid_pref'))}", color=GRAY)
    _para(doc, after=10)


def _annual_cash_flow(doc, sc, p1, p2):
    _page_break(doc)
    _heading(doc, 'Annual Cash Flow Summary')
    widths = [0.6, 0.95, 0.95, 0.6, 0.95, 0.95, 0.95, 0.95, 0.8, 0.8]
    t = _table(doc, 10, widths)
    _header_row(t, widths, ['Year', 'NOI', 'Debt Svc', 'DSCR', 'Levered CF', f'{p1} Dist',
                            f'{p2} Dist', 'Note Pmt', f'{p1} CoC', f'{p2} CoC'],
                first_align=WD_ALIGN_PARAGRAPH.CENTER)
    years = sc.get('annual_summary', [])
    tot = dict(noi=0.0, ds=0.0, cf=0.0, a=0.0, b=0.0, note=0.0)
    for i, yr in enumerate(years):
        fill = FILL_ALT if i % 2 == 0 else None
        dscr = yr.get('dscr') or 0.0
        vals = [str(yr.get('calendar_year')), fmt_k(yr.get('noi')), fmt_k(yr.get('debt_service')),
                f'{dscr:.2f}x', fmt_k(yr.get('levered_cf')), fmt_k(yr.get('distributions_ka')),
                fmt_k(yr.get('distributions_idp')), fmt_k(yr.get('surplus_note_payment')),
                fmt_pct(yr.get('coc_ka')), fmt_pct(yr.get('coc_idp'))]
        for j, (c, v) in enumerate(zip(_add_row(t, widths), vals)):
            color = None
            if j == 3:
                color = RED if dscr < 1.15 else (AMBER if dscr < 1.25 else None)
            _set_cell(c, v, align=WD_ALIGN_PARAGRAPH.RIGHT if j else WD_ALIGN_PARAGRAPH.CENTER,
                      fill=fill, color=color, bold=color is not None)
        tot['noi'] += yr.get('noi') or 0
        tot['ds'] += yr.get('debt_service') or 0
        tot['cf'] += yr.get('levered_cf') or 0
        tot['a'] += yr.get('distributions_ka') or 0
        tot['b'] += yr.get('distributions_idp') or 0
        tot['note'] += yr.get('surplus_note_payment') or 0
    totals = ['Total', fmt_k(tot['noi']), fmt_k(tot['ds']), '—', fmt_k(tot['cf']),
              fmt_k(tot['a']), fmt_k(tot['b']), fmt_k(tot['note']), '—', '—']
    for j, (c, v) in enumerate(zip(_add_row(t, widths), totals)):
        _set_cell(c, v, align=WD_ALIGN_PARAGRAPH.RIGHT if j else WD_ALIGN_PARAGRAPH.CENTER,
                  fill=FILL_TOTAL, bold=True)
    _para(doc, after=10)


def _debt_position(doc, sc):
    debt = sc['debt']
    _heading(doc, 'Debt Position')
    _kv_table(doc, [
        ('Current Loan Balance', fmt_usd(debt.get('current_balance'))),
        ('Original Principal', fmt_usd(debt.get('original_principal'))),
        ('Interest Rate', fmt_pct(debt.get('rate'))),
        ('Monthly Payment', fmt_usd(debt.get('monthly_payment'))),
        ('Annual Debt Service', fmt_usd(debt.get('annual_debt_service'))),
        ('Remaining Term', f"{debt.get('remaining_term_months')} months"),
        ('', ''),
        ('MIP Rate', fmt_pct(debt.get('mip_rate'))),
        ('Year 1 MIP', fmt_usd(debt.get('year1_mip'))),
        ('Total MIP (Hold Period)', fmt_usd(debt.get('total_mip_over_hold'))),
        ('', ''),
        ('Min DSCR', fmt_x(debt.get('min_dscr'))),
        ('Avg DSCR', fmt_x(debt.get('avg_dscr'))),
        ('Covenant Breaches', str(debt.get('breach_count'))),
        ('Initial LTV', fmt_pct(debt.get('initial_ltv'))),
        ('Terminal LTV', fmt_pct(debt.get('terminal_ltv'))),
    ])
    _para(doc, after=10)


def _scenario_comparison(doc, data, p1, p2):
    names = data.get('scenario_names') or []
    if len(names) < 2:
        return
    _page_break(doc)
    _heading(doc, 'TIF Scenario Comparison')
    _para(doc, 'Key return and risk metrics across all Tax Increment Financing scenarios. '
               'The first column represents the baseline (no appeal) case.',
          color=GRAY, italic=True, after=6)
    metrics = [('Deal IRR', 'deal_irr', fmt_pct), ('Deal EM', 'deal_em', fmt_x),
               (f'{p1} EM', 'sponsor_em', fmt_x), (f'{p2} EM', 'investor_em', fmt_x),
               ('Min DSCR', 'min_dscr', fmt_x), ('Avg DSCR', 'avg_dscr', fmt_x),
               ('Total Distributions', 'total_distributions', fmt_usd),
               ('Net Sale Proceeds', 'net_sale_proceeds', fmt_usd)]
    label_w = 2.0
    col_w = (8.8 - label_w) / len(names)
    widths = [label_w] + [col_w] * len(names)
    t = _table(doc, len(widths), widths)
    _header_row(t, widths, ['Metric'] + list(names))
    for mi, (label, key, fn) in enumerate(metrics):
        fill = FILL_ALT if mi % 2 == 0 else None
        cells = _add_row(t, widths)
        _set_cell(cells[0], label, fill=fill)
        for i, n in enumerate(names):
            dm = (data['scenarios'].get(n) or {}).get('decision_metrics') or {}
            _set_cell(cells[i + 1], fn(dm.get(key)), align=WD_ALIGN_PARAGRAPH.RIGHT,
                      fill=fill, bold=(i == 0))
    _para(doc, after=10)


def _disclaimer(doc):
    p = _para(doc, after=0)
    p.paragraph_format.space_before = Pt(20)
    _run(p, 'DISCLAIMER: This report is generated from modeled projections and is intended for '
            'informational purposes only. Actual results may differ materially from projections. '
            'This does not constitute investment advice. Consult with qualified legal, tax, and '
            'financial advisors before making investment decisions.',
         size=8, color=GRAY, italic=True)


# ─── entry point ───────────────────────────────────────────────────

def build_investor_report(data: dict, out_path: str) -> str:
    names = data.get('scenario_names') or []
    primary = data.get('_primary_scenario') or (names[0] if names else None)
    sc = (data.get('scenarios') or {}).get(primary)
    if not sc:
        raise ValueError(f'Scenario "{primary}" not found')
    ids = data.get('partner_ids') or []
    p1 = ids[0] if len(ids) > 0 and ids[0] else 'P1'
    p2 = ids[1] if len(ids) > 1 and ids[1] else 'P2'
    today = datetime.date.today()
    date_str = f'{today:%B} {today.day}, {today.year}'

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(11)

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for side in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(sec, side, Inches(0.75))

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(hp, f"{data.get('entity_name') or 'Partnership'} — Investor Report", size=8, color=GRAY)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(fp, 'Page ', size=8, color=GRAY)
    _page_number_field(fp)
    _run(fp, f'  |  Generated {date_str}', size=8, color=GRAY)

    _cover(doc, data, sc, date_str)
    _deal_summary(doc, sc)
    _partner_returns(doc, sc)
    _annual_cash_flow(doc, sc, p1, p2)
    _debt_position(doc, sc)
    _scenario_comparison(doc, data, p1, p2)
    _disclaimer(doc)
    doc.save(out_path)
    return out_path
